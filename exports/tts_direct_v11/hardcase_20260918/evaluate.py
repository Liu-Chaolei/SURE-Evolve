"""Independent frozen-artifact Seed-TTS evaluation; never writes training state."""
import concurrent.futures
import fcntl
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import time

ROOT = Path(__file__).resolve().parent
PLAN = json.loads((ROOT / 'plan.json').read_text())
PROJECT = ROOT.parents[2]
OUTPUT = ROOT.parent / 'TTS-Direct-v11_Seed-TTS中文_测试结果_20260918.xlsx'


def write(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix('.pending')
    temp.write_text(json.dumps(data, ensure_ascii=False, indent=2))
    temp.replace(path)


def worker(index):
    row = PLAN['models'][index]
    work = Path(row['workspace'])
    status = work / 'status.json'
    try:
        write(status, {'stage': 'verify', 'started_at': time.time()})
        manifest = Path(row['manifest'])
        assert hashlib.sha256(manifest.read_bytes()).hexdigest() == row['manifest_sha256']
        bundle = json.loads(manifest.read_text())
        # Check every frozen asset, including weights, source and training evidence.
        for relative, expected in bundle['files'].items():
            h = hashlib.sha256()
            with (manifest.parent / relative).open('rb') as f:
                for data in iter(lambda: f.read(8*1024*1024), b''):
                    h.update(data)
            if h.hexdigest() != expected:
                raise ValueError('Artifact hash mismatch: ' + relative)
        env = dict(os.environ, LOCAL_RANK=str(index), SURE_ACCELERATOR='npu',
                   SURE_CPU_THREADS='8', OMP_NUM_THREADS='8', MKL_NUM_THREADS='8',
                   SURE_PRECISION='fp32', SURE_TTS_TRAIN_SEED='42',
                   TOKENIZERS_PARALLELISM='false', HF_HUB_OFFLINE='1',
                   PYTHONPATH=PLAN['source'], PYTHONDONTWRITEBYTECODE='1')
        tools = Path(PLAN['source']) / 'playground/sure_master/tools'
        r = row['resources']
        cmd = [sys.executable, str(tools/'run_f5tts_batch_infer.py'),
               '--f5-root', r['source'], '--eval-data', PLAN['holdout'],
               '--language', 'zh', '--max-samples', '0', '--workers', '1',
               '--device', f'npu:{index}', '--ckpt-file', r['checkpoint'],
               '--vocab-file', r['vocab'], '--model-cfg', r['model_cfg'],
               '--load-vocoder-from-local', '1', '--vocoder-local-path', r['vocoder'],
               '--candidate-type', 'inference', '--resume', '1', '--timeout', '0']
        for k,v in row['inference_config'].items():
            cmd.extend(['--'+k.replace('_','-'),str(int(v)) if isinstance(v,bool) else str(v)])
        write(work/'inference_command.json', cmd)
        write(status, {'stage': 'synthesis', 'started_at': time.time(), 'device': index})
        with (work/'synthesis.log').open('a') as log:
            subprocess.run(cmd, cwd=work, env=env, stdout=log, stderr=subprocess.STDOUT, check=True)
        sys.path.insert(0, PLAN['source'])
        from playground.sure_master.tasks.tts_outputs import validate_samples
        validate_samples(Path(PLAN['holdout']), work/'artifacts/samples.jsonl')
        write(status, {'stage': 'scoring', 'updated_at': time.time()})
        with (work/'scoring.log').open('a') as log:
            subprocess.run([sys.executable,str(tools/'score_task.py'),str(work/'score_request.json'),str(work/'score_result.json')],cwd=work,env=env,stdout=log,stderr=subprocess.STDOUT,check=True)
        result=json.loads((work/'score_result.json').read_text())
        if not result.get('success'):
            raise RuntimeError(str(result))
        write(status, {'stage': 'completed', 'score': result['score'], 'finished_at': time.time()})
        return True
    except Exception as e:
        write(status, {'stage': 'failed', 'error': str(e), 'finished_at': time.time()})
        return False


def export():
    from copy import copy
    from openpyxl import load_workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo
    template = PROJECT/'exports/asr_direct_v1/ASR-Direct-v1_TEDLIUM3_测试结果_20260917.xlsx'
    wb=load_workbook(template)
    names=['结果概览','数据与实验设置','搜索阶段参考','全部模型测试','各轮事后汇总','产物溯源']
    for ws,name in zip(wb.worksheets,names):
        ws.title=name
        for key in list(ws.tables): del ws.tables[key]
        for row in ws:
            for c in row: c.value=None
    measured=[]
    for row in PLAN['models']:
        work=Path(row['workspace']); p=work/'metric/report.json'
        status=json.loads((work/'status.json').read_text()) if (work/'status.json').exists() else {'stage':'pending'}
        report=json.loads(p.read_text()) if p.exists() and status['stage']=='completed' else {}
        score=report.get('score')
        counts=report.get('details',{}).get('scoring_result',{})
        if score is not None and counts.get('num_ref_utts') != PLAN['sample_count']:
            raise ValueError('Report sample count mismatch for '+row['name'])
        measured.append({**row,'score':score,'counts':counts,'stage':status['stage'],'error':status.get('error','')})
    baseline=measured[0]['score']; completed=[r for r in measured if r['score'] is not None]
    candidates=[r for r in completed if r['round']>0]
    best=min(candidates,key=lambda r:r['score']) if candidates else None
    data=[]
    data.append([['项目','内容'],['部署','TTS-Direct-v11'],['更新时间',time.strftime('%Y-%m-%d %H:%M:%S %Z')],['测试规模',f'baseline＋7个候选，{len(completed)}/8完成'],['基线 holdout CER',baseline],['事后补测最优',best['name'] if best else '待完成'],['事后最优 CER',best['score'] if best else None],['范围','截至启动时7个已封存候选；exp_8/exp_10正在重跑，不纳入'],['结果解释','中途独立holdout补测，不回写搜索历史，不替换正式胜者'],['统计范围','seed42单次推理；仅本批冻结模型；未做显著性检验'],['指标口径','SURE Paraformer中文转写＋去标点＋WeNet corpus CER；非官方逐句平均WER'],['当前状态','全部完成' if len(completed)==8 else '测试中；空白CER表示尚未完成']])
    data.append([['项目','设置','说明'],['测试集','Seed-TTS 中文', '完整zh/meta.lst对应冻结holdout'],['样本数',PLAN['sample_count'],'全部模型一致'],['训练数据','WenetSpeech4TTS-Premium','本次不训练'],['模型','F5-TTS v1 Base及候选','保留结构与源代码'],['训练完成','50 epochs','final EMA'],['推理精度','FP32','与旧部署搜索一致'],['推理seed',42,'每模型1 worker'],['模型并发',8,'n19 / 10052 / 每模型1 NPU'],['推理参数','各候选冻结inference_config','包括nfe、CFG、cross_fade等'],['CER','(替换＋删除＋插入)/参考字符数','corpus_edit_distance'],['评分','Paraformer CPU，8线程/模型','与旧部署搜索评分管线相同'],['manifest SHA256',PLAN['holdout_sha256'],'独立副本'],['表格模板',str(template),'保留ASR原文件']])
    data.append([['模型','Search CER','Selection CER','Holdout CER','绝对下降（百分点）','相对CER下降','说明']]+[[r['name'],r['search_cer'],None,r['score'],(baseline-r['score'])*100 if baseline is not None and r['score'] is not None else None,(baseline-r['score'])/baseline if baseline and r['score'] is not None else None,'独立补测；尚无正式最终选模结果'] for r in measured])
    ordered=sorted(measured,key=lambda r:(r['score'] is None,r['score'] or 0))
    data.append([['模型','轮次（0=基线）','Search CER','Selection CER','Holdout CER','绝对下降（百分点）','相对CER下降','错误总数','替换','删除','插入','参考字符数','对比基线','状态','测试来源']])
    for r in ordered:
        c=r['counts']; s=r['score']; delta=baseline-s if baseline is not None and s is not None else None
        data[3].append([r['name'],r['round'],r['search_cer'],None,s,delta*100 if delta is not None else None,delta/baseline if delta is not None and baseline else None,sum(c.get(k,0) for k in ['sub','del','ins']) if c else None,c.get('sub'),c.get('del'),c.get('ins'),c.get('all'),('优于基线' if delta>0 else '退化' if delta<0 else '持平') if delta is not None else None,r['stage']+(': '+r['error'] if r['error'] else ''),'Seed-TTS中文独立补测'])
    data.append([['轮次','候选数','本轮Holdout最低模型','本轮最低CER','本轮平均CER','优于基线数','持平数','退化数','本轮最低相对改善']])
    for rnd in [1,2,3]:
        rr=[r for r in candidates if r['round']==rnd]; b=min(rr,key=lambda r:r['score']) if rr else None
        data[4].append([rnd,len(rr),b['name'] if b else None,b['score'] if b else None,sum(r['score'] for r in rr)/len(rr) if rr else None,*([sum(r['score']<baseline for r in rr),sum(r['score']==baseline for r in rr),sum(r['score']>baseline for r in rr)] if baseline is not None else [None]*3),(baseline-b['score'])/baseline if baseline and b else None])
    data.append([['模型','Idea ID','模型manifest','SURE评测报告','测试工作目录']]+[[r['name'],r['idea_id'],r['manifest'],str(Path(r['workspace'])/'metric/report.json') if r['score'] is not None else None,r['workspace']] for r in ordered])
    for i,(ws,rows) in enumerate(zip(wb.worksheets,data)):
        for ri,vals in enumerate(rows,1):
            for ci,val in enumerate(vals,1):
                c=ws.cell(ri,ci,val)
                if ri>2: c._style=copy(ws.cell(2,ci)._style)
        if ws.max_row>len(rows): ws.delete_rows(len(rows)+1,ws.max_row-len(rows))
        if i in [2,3]:
            for ri in range(2,len(rows)+1):
                for ci in ([2,3,4,6] if i==2 else [3,4,5,7]): ws.cell(ri,ci).number_format='0.0000%'
        tab=Table(displayName=f'TTSResults{i+1}',ref=f'A1:{ws.cell(len(rows),len(rows[0])).coordinate}')
        tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(tab)
    tmp=OUTPUT.with_suffix('.pending.xlsx');wb.save(tmp);tmp.replace(OUTPUT)
    write(ROOT/'aggregate.json',measured)


def host():
    # Cooperate with the main controller's existing allocation lock.
    lock=PROJECT/'runs/f5tts_allocations_10024/10052.lock'
    with lock.open('a') as handle:
        fcntl.flock(handle,fcntl.LOCK_EX|fcntl.LOCK_NB)
        cmd=['srun','--jobid=10052','--overlap','--exact','--ntasks=1','--cpus-per-task=64','--gres=gpu:ascend910b3:8','sudo','-n','slurm-docker-run','--pull','missing','--network','host','--shm-size','64g','--mount','/shared/chaolei.liu:/shared/chaolei.liu:ro','--mount',f'{ROOT}:{ROOT}:rw','--workdir',str(ROOT),'--env','PYTHONDONTWRITEBYTECODE=1',PLAN['image'],'python',str(Path(__file__).resolve()),'pool']
        write(ROOT/'launch_command.json',cmd)
        with (ROOT/'pool.log').open('a') as log:
            proc=subprocess.Popen(cmd,stdout=log,stderr=subprocess.STDOUT)
            write(ROOT/'driver.json',{'pid':os.getpid(),'srun_pid':proc.pid,'job_id':'10052','started_at':time.time(),'status':'running'})
            while proc.poll() is None:
                export();time.sleep(30)
            export()
            write(ROOT/'driver.json',{'pid':os.getpid(),'job_id':'10052','exit_code':proc.returncode,'status':'completed' if proc.returncode==0 else 'failed','finished_at':time.time()})

if __name__=='__main__':
    if sys.argv[1]=='pool':
        with concurrent.futures.ThreadPoolExecutor(max_workers=8) as pool:
            results=list(pool.map(worker,range(len(PLAN['models']))))
        sys.exit(0 if all(results) else 1)
    elif sys.argv[1]=='export': export()
    else: host()
